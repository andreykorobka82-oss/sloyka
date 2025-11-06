from fastapi import FastAPI, APIRouter, HTTPException, Depends
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional
import uuid
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from fastapi.responses import StreamingResponse
import io

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# Create the main app without a prefix
app = FastAPI()

# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")

# ============ Models ============

class User(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    pin: str
    role: str = "user"  # admin or user
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class UserCreate(BaseModel):
    name: str
    pin: str
    role: str = "user"

class LoginRequest(BaseModel):
    pin: str

class Category(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class CategoryCreate(BaseModel):
    name: str

class Product(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    category_id: str
    price: float
    current_stock: float = 0
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class ProductCreate(BaseModel):
    name: str
    category_id: str
    price: float

class ProductUpdate(BaseModel):
    name: Optional[str] = None
    category_id: Optional[str] = None
    price: Optional[float] = None

class Income(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    product_id: str
    quantity: float
    user_id: str
    user_name: str
    product_name: str
    date: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class IncomeCreate(BaseModel):
    product_id: str
    quantity: float
    user_id: str

class IncomeUpdate(BaseModel):
    quantity: Optional[float] = None

class Expense(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    product_id: str
    quantity: float
    user_id: str
    user_name: str
    product_name: str
    date: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class ExpenseCreate(BaseModel):
    product_id: str
    quantity: float
    user_id: str

class ExpenseUpdate(BaseModel):
    quantity: Optional[float] = None

class Revenue(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    date: datetime
    amount: float
    user_id: str
    user_name: str
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class RevenueCreate(BaseModel):
    date: datetime
    amount: float
    user_id: str

class RevenueUpdate(BaseModel):
    date: Optional[datetime] = None
    amount: Optional[float] = None

class RecountRequest(BaseModel):
    start_date: datetime
    end_date: datetime
    final_stocks: dict  # product_id -> final_stock

class RecountProductData(BaseModel):
    product_id: str
    product_name: str
    category_name: str
    initial_stock: float
    final_stock: float
    difference: float
    price: float
    sale_amount: float

class RecountResult(BaseModel):
    start_date: datetime
    end_date: datetime
    total_revenue: float
    products: List[RecountProductData]
    total_sales: float
    result: float
    categories_sales: dict

# ============ Helper Functions ============

async def get_user_by_id(user_id: str):
    user = await db.users.find_one({"id": user_id}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return user

async def get_product_by_id(product_id: str):
    product = await db.products.find_one({"id": product_id}, {"_id": 0})
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    return product

async def initialize_admin():
    # Check if admin already exists
    admin = await db.users.find_one({"pin": "5534"}, {"_id": 0})
    if not admin:
        admin_user = User(
            name="Адміністратор",
            pin="5534",
            role="admin"
        )
        doc = admin_user.model_dump()
        doc['created_at'] = doc['created_at'].isoformat()
        await db.users.insert_one(doc)
        logging.info("Default admin created")

# ============ Auth Routes ============

@api_router.post("/auth/login")
async def login(request: LoginRequest):
    user = await db.users.find_one({"pin": request.pin}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=401, detail="Невірний ПІН-код")
    
    if isinstance(user['created_at'], str):
        user['created_at'] = datetime.fromisoformat(user['created_at'])
    
    return {"user": user, "message": "Вхід успішний"}

# ============ User Routes ============

@api_router.get("/users", response_model=List[User])
async def get_users():
    users = await db.users.find({}, {"_id": 0}).to_list(1000)
    for user in users:
        if isinstance(user['created_at'], str):
            user['created_at'] = datetime.fromisoformat(user['created_at'])
    return users

@api_router.post("/users", response_model=User)
async def create_user(input: UserCreate):
    # Check if PIN already exists
    existing = await db.users.find_one({"pin": input.pin}, {"_id": 0})
    if existing:
        raise HTTPException(status_code=400, detail="ПІН-код вже використовується")
    
    user = User(**input.model_dump())
    doc = user.model_dump()
    doc['created_at'] = doc['created_at'].isoformat()
    await db.users.insert_one(doc)
    return user

@api_router.delete("/users/{user_id}")
async def delete_user(user_id: str):
    result = await db.users.delete_one({"id": user_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Користувача не знайдено")
    return {"message": "Користувача видалено"}

# ============ Category Routes ============

@api_router.get("/categories", response_model=List[Category])
async def get_categories():
    categories = await db.categories.find({}, {"_id": 0}).to_list(1000)
    for cat in categories:
        if isinstance(cat['created_at'], str):
            cat['created_at'] = datetime.fromisoformat(cat['created_at'])
    return categories

@api_router.post("/categories", response_model=Category)
async def create_category(input: CategoryCreate):
    category = Category(**input.model_dump())
    doc = category.model_dump()
    doc['created_at'] = doc['created_at'].isoformat()
    await db.categories.insert_one(doc)
    return category

@api_router.put("/categories/{category_id}", response_model=Category)
async def update_category(category_id: str, input: CategoryCreate):
    result = await db.categories.update_one(
        {"id": category_id},
        {"$set": {"name": input.name}}
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Категорію не знайдено")
    
    category = await db.categories.find_one({"id": category_id}, {"_id": 0})
    if isinstance(category['created_at'], str):
        category['created_at'] = datetime.fromisoformat(category['created_at'])
    return category

@api_router.delete("/categories/{category_id}")
async def delete_category(category_id: str):
    result = await db.categories.delete_one({"id": category_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Категорію не знайдено")
    return {"message": "Категорію видалено"}

# ============ Product Routes ============

@api_router.get("/products", response_model=List[Product])
async def get_products():
    products = await db.products.find({}, {"_id": 0}).to_list(1000)
    for product in products:
        if isinstance(product['created_at'], str):
            product['created_at'] = datetime.fromisoformat(product['created_at'])
    return products

@api_router.get("/products/in-stock", response_model=List[Product])
async def get_products_in_stock():
    products = await db.products.find({"current_stock": {"$gt": 0}}, {"_id": 0}).to_list(1000)
    for product in products:
        if isinstance(product['created_at'], str):
            product['created_at'] = datetime.fromisoformat(product['created_at'])
    return products

@api_router.post("/products", response_model=Product)
async def create_product(input: ProductCreate):
    product = Product(**input.model_dump())
    doc = product.model_dump()
    doc['created_at'] = doc['created_at'].isoformat()
    await db.products.insert_one(doc)
    return product

@api_router.put("/products/{product_id}", response_model=Product)
async def update_product(product_id: str, input: ProductUpdate):
    update_data = {k: v for k, v in input.model_dump().items() if v is not None}
    if not update_data:
        raise HTTPException(status_code=400, detail="Немає даних для оновлення")
    
    result = await db.products.update_one(
        {"id": product_id},
        {"$set": update_data}
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Товар не знайдено")
    
    product = await db.products.find_one({"id": product_id}, {"_id": 0})
    if isinstance(product['created_at'], str):
        product['created_at'] = datetime.fromisoformat(product['created_at'])
    return product

@api_router.delete("/products/{product_id}")
async def delete_product(product_id: str):
    result = await db.products.delete_one({"id": product_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Товар не знайдено")
    return {"message": "Товар видалено"}

# ============ Income Routes ============

@api_router.get("/incomes", response_model=List[Income])
async def get_incomes():
    incomes = await db.incomes.find({}, {"_id": 0}).sort("date", -1).to_list(1000)
    for income in incomes:
        if isinstance(income['date'], str):
            income['date'] = datetime.fromisoformat(income['date'])
        if isinstance(income['created_at'], str):
            income['created_at'] = datetime.fromisoformat(income['created_at'])
    return incomes

@api_router.post("/incomes", response_model=Income)
async def create_income(input: IncomeCreate):
    # Get user and product info
    user = await get_user_by_id(input.user_id)
    product = await get_product_by_id(input.product_id)
    
    income = Income(
        product_id=input.product_id,
        quantity=input.quantity,
        user_id=input.user_id,
        user_name=user['name'],
        product_name=product['name']
    )
    
    # Update product stock
    await db.products.update_one(
        {"id": input.product_id},
        {"$inc": {"current_stock": input.quantity}}
    )
    
    doc = income.model_dump()
    doc['date'] = doc['date'].isoformat()
    doc['created_at'] = doc['created_at'].isoformat()
    await db.incomes.insert_one(doc)
    return income

@api_router.put("/incomes/{income_id}", response_model=Income)
async def update_income(income_id: str, input: IncomeUpdate):
    # Get existing income
    existing_income = await db.incomes.find_one({"id": income_id}, {"_id": 0})
    if not existing_income:
        raise HTTPException(status_code=404, detail="Надходження не знайдено")
    
    # Calculate stock difference
    old_quantity = existing_income['quantity']
    new_quantity = input.quantity if input.quantity is not None else old_quantity
    stock_diff = new_quantity - old_quantity
    
    # Update product stock
    await db.products.update_one(
        {"id": existing_income['product_id']},
        {"$inc": {"current_stock": stock_diff}}
    )
    
    # Update income
    await db.incomes.update_one(
        {"id": income_id},
        {"$set": {"quantity": new_quantity}}
    )
    
    income = await db.incomes.find_one({"id": income_id}, {"_id": 0})
    if isinstance(income['date'], str):
        income['date'] = datetime.fromisoformat(income['date'])
    if isinstance(income['created_at'], str):
        income['created_at'] = datetime.fromisoformat(income['created_at'])
    return income

@api_router.delete("/incomes/{income_id}")
async def delete_income(income_id: str):
    # Get income to revert stock
    income = await db.incomes.find_one({"id": income_id}, {"_id": 0})
    if not income:
        raise HTTPException(status_code=404, detail="Надходження не знайдено")
    
    # Revert stock
    await db.products.update_one(
        {"id": income['product_id']},
        {"$inc": {"current_stock": -income['quantity']}}
    )
    
    await db.incomes.delete_one({"id": income_id})
    return {"message": "Надходження видалено"}

# ============ Expense Routes ============

@api_router.get("/expenses", response_model=List[Expense])
async def get_expenses():
    expenses = await db.expenses.find({}, {"_id": 0}).sort("date", -1).to_list(1000)
    for expense in expenses:
        if isinstance(expense['date'], str):
            expense['date'] = datetime.fromisoformat(expense['date'])
        if isinstance(expense['created_at'], str):
            expense['created_at'] = datetime.fromisoformat(expense['created_at'])
    return expenses

@api_router.post("/expenses", response_model=Expense)
async def create_expense(input: ExpenseCreate):
    # Get user and product info
    user = await get_user_by_id(input.user_id)
    product = await get_product_by_id(input.product_id)
    
    # Check if enough stock
    if product['current_stock'] < input.quantity:
        raise HTTPException(status_code=400, detail="Недостатньо товару на складі")
    
    expense = Expense(
        product_id=input.product_id,
        quantity=input.quantity,
        user_id=input.user_id,
        user_name=user['name'],
        product_name=product['name']
    )
    
    # Update product stock
    await db.products.update_one(
        {"id": input.product_id},
        {"$inc": {"current_stock": -input.quantity}}
    )
    
    doc = expense.model_dump()
    doc['date'] = doc['date'].isoformat()
    doc['created_at'] = doc['created_at'].isoformat()
    await db.expenses.insert_one(doc)
    return expense

@api_router.put("/expenses/{expense_id}", response_model=Expense)
async def update_expense(expense_id: str, input: ExpenseUpdate):
    # Get existing expense
    existing_expense = await db.expenses.find_one({"id": expense_id}, {"_id": 0})
    if not existing_expense:
        raise HTTPException(status_code=404, detail="Списання не знайдено")
    
    # Calculate stock difference
    old_quantity = existing_expense['quantity']
    new_quantity = input.quantity if input.quantity is not None else old_quantity
    stock_diff = old_quantity - new_quantity  # Reverse logic for expenses
    
    # Update product stock
    await db.products.update_one(
        {"id": existing_expense['product_id']},
        {"$inc": {"current_stock": stock_diff}}
    )
    
    # Update expense
    await db.expenses.update_one(
        {"id": expense_id},
        {"$set": {"quantity": new_quantity}}
    )
    
    expense = await db.expenses.find_one({"id": expense_id}, {"_id": 0})
    if isinstance(expense['date'], str):
        expense['date'] = datetime.fromisoformat(expense['date'])
    if isinstance(expense['created_at'], str):
        expense['created_at'] = datetime.fromisoformat(expense['created_at'])
    return expense

@api_router.delete("/expenses/{expense_id}")
async def delete_expense(expense_id: str):
    # Get expense to revert stock
    expense = await db.expenses.find_one({"id": expense_id}, {"_id": 0})
    if not expense:
        raise HTTPException(status_code=404, detail="Списання не знайдено")
    
    # Revert stock
    await db.products.update_one(
        {"id": expense['product_id']},
        {"$inc": {"current_stock": expense['quantity']}}
    )
    
    await db.expenses.delete_one({"id": expense_id})
    return {"message": "Списання видалено"}

# ============ Revenue Routes ============

@api_router.get("/revenues", response_model=List[Revenue])
async def get_revenues():
    revenues = await db.revenues.find({}, {"_id": 0}).sort("date", -1).to_list(1000)
    for revenue in revenues:
        if isinstance(revenue['date'], str):
            revenue['date'] = datetime.fromisoformat(revenue['date'])
        if isinstance(revenue['created_at'], str):
            revenue['created_at'] = datetime.fromisoformat(revenue['created_at'])
    return revenues

@api_router.post("/revenues", response_model=Revenue)
async def create_revenue(input: RevenueCreate):
    user = await get_user_by_id(input.user_id)
    
    revenue = Revenue(
        date=input.date,
        amount=input.amount,
        user_id=input.user_id,
        user_name=user['name']
    )
    
    doc = revenue.model_dump()
    doc['date'] = doc['date'].isoformat()
    doc['created_at'] = doc['created_at'].isoformat()
    await db.revenues.insert_one(doc)
    return revenue

@api_router.put("/revenues/{revenue_id}", response_model=Revenue)
async def update_revenue(revenue_id: str, input: RevenueUpdate):
    update_data = {}
    if input.date is not None:
        update_data['date'] = input.date.isoformat()
    if input.amount is not None:
        update_data['amount'] = input.amount
    
    if not update_data:
        raise HTTPException(status_code=400, detail="Немає даних для оновлення")
    
    result = await db.revenues.update_one(
        {"id": revenue_id},
        {"$set": update_data}
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Виручку не знайдено")
    
    revenue = await db.revenues.find_one({"id": revenue_id}, {"_id": 0})
    if isinstance(revenue['date'], str):
        revenue['date'] = datetime.fromisoformat(revenue['date'])
    if isinstance(revenue['created_at'], str):
        revenue['created_at'] = datetime.fromisoformat(revenue['created_at'])
    return revenue

@api_router.delete("/revenues/{revenue_id}")
async def delete_revenue(revenue_id: str):
    result = await db.revenues.delete_one({"id": revenue_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Виручку не знайдено")
    return {"message": "Виручку видалено"}

# ============ Recount Routes ============

@api_router.post("/recount/generate", response_model=RecountResult)
async def generate_recount(request: RecountRequest):
    # Get all products
    products = await db.products.find({}, {"_id": 0}).to_list(1000)
    categories = await db.categories.find({}, {"_id": 0}).to_list(1000)
    categories_dict = {cat['id']: cat['name'] for cat in categories}
    
    # Calculate total revenue for the period
    revenues = await db.revenues.find({}, {"_id": 0}).to_list(1000)
    total_revenue = sum(
        rev['amount'] for rev in revenues
        if isinstance(rev['date'], str) and 
           request.start_date <= datetime.fromisoformat(rev['date']) <= request.end_date
    )
    
    # Get incomes and expenses for the period
    incomes = await db.incomes.find({}, {"_id": 0}).to_list(1000)
    expenses = await db.expenses.find({}, {"_id": 0}).to_list(1000)
    
    # Calculate for each product
    product_data_list = []
    categories_sales = {}
    
    for product in products:
        product_id = product['id']
        
        # Calculate initial stock (current - incomes + expenses during period)
        period_incomes = sum(
            inc['quantity'] for inc in incomes
            if inc['product_id'] == product_id and
               isinstance(inc['date'], str) and
               request.start_date <= datetime.fromisoformat(inc['date']) <= request.end_date
        )
        
        period_expenses = sum(
            exp['quantity'] for exp in expenses
            if exp['product_id'] == product_id and
               isinstance(exp['date'], str) and
               request.start_date <= datetime.fromisoformat(exp['date']) <= request.end_date
        )
        
        initial_stock = product['current_stock'] - period_incomes + period_expenses
        final_stock = request.final_stocks.get(product_id, product['current_stock'])
        difference = final_stock - initial_stock
        sale_amount = abs(difference) * product['price'] if difference < 0 else 0
        
        category_name = categories_dict.get(product['category_id'], 'Без категорії')
        
        product_data = RecountProductData(
            product_id=product_id,
            product_name=product['name'],
            category_name=category_name,
            initial_stock=initial_stock,
            final_stock=final_stock,
            difference=difference,
            price=product['price'],
            sale_amount=sale_amount
        )
        product_data_list.append(product_data)
        
        # Accumulate category sales
        if category_name not in categories_sales:
            categories_sales[category_name] = 0
        categories_sales[category_name] += sale_amount
    
    total_sales = sum(pd.sale_amount for pd in product_data_list)
    result = total_revenue - total_sales
    
    return RecountResult(
        start_date=request.start_date,
        end_date=request.end_date,
        total_revenue=total_revenue,
        products=product_data_list,
        total_sales=total_sales,
        result=result,
        categories_sales=categories_sales
    )

@api_router.post("/recount/export")
async def export_recount(request: RecountRequest):
    # Generate recount data
    recount_data = await generate_recount(request)
    
    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Переоблік"
    
    # Set column widths
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 15
    
    # Header style
    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="04A4F4", end_color="04A4F4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    # Add title
    ws['A1'] = 'Звіт Переобліку'
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:G1')
    
    # Add period
    ws['A2'] = f"Період: {recount_data.start_date.strftime('%d.%m.%Y')} - {recount_data.end_date.strftime('%d.%m.%Y')}"
    ws.merge_cells('A2:G2')
    
    # Add headers
    row = 4
    headers = ['Товар', 'Категорія', 'Початковий залишок', 'Кінцевий залишок', 'Різниця', 'Ціна', 'Сума продажу']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # Add product data
    row = 5
    for product in recount_data.products:
        ws.cell(row=row, column=1, value=product.product_name)
        ws.cell(row=row, column=2, value=product.category_name)
        ws.cell(row=row, column=3, value=product.initial_stock)
        ws.cell(row=row, column=4, value=product.final_stock)
        ws.cell(row=row, column=5, value=product.difference)
        ws.cell(row=row, column=6, value=product.price)
        ws.cell(row=row, column=7, value=product.sale_amount)
        row += 1
    
    # Add summary
    row += 1
    ws.cell(row=row, column=1, value="Загальна сума виручок:").font = Font(bold=True)
    ws.cell(row=row, column=7, value=recount_data.total_revenue)
    
    row += 1
    ws.cell(row=row, column=1, value="Загальна сума продажів:").font = Font(bold=True)
    ws.cell(row=row, column=7, value=recount_data.total_sales)
    
    row += 1
    ws.cell(row=row, column=1, value="Результат переобліку:").font = Font(bold=True)
    result_cell = ws.cell(row=row, column=7, value=recount_data.result)
    result_cell.font = Font(bold=True, color="E53E3E" if recount_data.result < 0 else "04A4F4")
    
    # Add category sales
    row += 2
    ws.cell(row=row, column=1, value="Продажі по категоріям:").font = Font(bold=True)
    row += 1
    for category, amount in recount_data.categories_sales.items():
        ws.cell(row=row, column=1, value=category)
        ws.cell(row=row, column=7, value=amount)
        row += 1
    
    # Save to BytesIO
    excel_file = io.BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)
    
    # Return as streaming response
    filename = f"recount_{recount_data.start_date.strftime('%Y%m%d')}_{recount_data.end_date.strftime('%Y%m%d')}.xlsx"
    
    return StreamingResponse(
        excel_file,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# Include the router in the main app
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("startup")
async def startup_event():
    await initialize_admin()
    logger.info("Application started")

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()